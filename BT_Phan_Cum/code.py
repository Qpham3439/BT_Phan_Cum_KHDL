# ==============================================================================
# IMPORT THƯ VIỆN
# ==============================================================================

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

import openpyxl

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)

from sklearn.cluster import KMeans

# ==============================================================================
# 1. ĐỌC FILE CSV
# ==============================================================================

file_path = "TỔNG HỢP ĐIỂM K58KTP - Trang tính1.csv"

df_raw = pd.read_csv(
    file_path,
    header=None,
    encoding_errors='ignore',
    engine='python'
)

print("Kích thước dữ liệu:")
print(df_raw.shape)

print("\nXem trước dữ liệu:")
print(df_raw.head(10))

# ==============================================================================
# 2. LẤY MSSV VÀ TÊN SINH VIÊN
# ==============================================================================

# File CSV của m:
# dòng 1 = MSSV
# dòng 2 = Tên sinh viên

mssv_row = df_raw.iloc[1, 3:].values

tensv_row = df_raw.iloc[2, 3:].values

# ==============================================================================
# 3. LẤY MA TRẬN ĐIỂM
# ==============================================================================

# bắt đầu lấy điểm từ dòng 4 trở đi

students_df = df_raw.iloc[4:, 3:].copy()

# đổi dấu phẩy thành dấu chấm
students_df = students_df.replace(',', '.', regex=True)

# chuyển sang số
students_df = students_df.apply(
    pd.to_numeric,
    errors='coerce'
)

print("\nMa trận điểm:")
print(students_df.head())

# ==============================================================================
# 4. XÓA CỘT TRỐNG TOÀN BỘ
# ==============================================================================

cols_to_drop = students_df.columns[
    students_df.isna().all(axis=0)
]

students_df_filtered = students_df.drop(
    columns=cols_to_drop
)

# xóa MSSV/Tên tương ứng cột bị drop
mssv_filtered = np.delete(
    mssv_row,
    [c - 3 for c in cols_to_drop]
)

tensv_filtered = np.delete(
    tensv_row,
    [c - 3 for c in cols_to_drop]
)

print("\nSố sinh viên hợp lệ:")
print(len(mssv_filtered))

# ==============================================================================
# 5. XỬ LÝ Ô TRỐNG
# ==============================================================================

students_df_final = students_df_filtered.fillna(0.0)

# ==============================================================================
# 6. TÍNH GPA
# ==============================================================================

gpa_scores = students_df_final.mean(axis=0)

features_filtered = pd.DataFrame({
    'MSSV': mssv_filtered,
    'Ten_SV': tensv_filtered,
    'GPA': gpa_scores.values
})

print("\nDanh sách GPA:")
print(features_filtered.head())

# ==============================================================================
# 7. ELBOW METHOD
# ==============================================================================

X = features_filtered[['GPA']].values

inertia = []

K = range(1, 10)

for k in K:

    model = KMeans(
        n_clusters=k,
        random_state=42,
        n_init=10
    )

    model.fit(X)

    inertia.append(model.inertia_)

plt.figure(figsize=(8,5))

plt.plot(K, inertia, marker='o')

plt.title("Elbow Method")

plt.xlabel("So cum K")

plt.ylabel("Inertia")

plt.grid(True)

plt.show()

# ==============================================================================
# 8. PHÂN CỤM KMEANS
# ==============================================================================

kmeans = KMeans(
    n_clusters=3,
    random_state=42,
    n_init=10
)

features_filtered['Cluster_Label'] = (
    kmeans.fit_predict(X)
)

print("\nTâm cụm:")
print(kmeans.cluster_centers_)

# ==============================================================================
# 9. ĐẶT TÊN NHÓM
# ==============================================================================

cluster_means = features_filtered.groupby(
    'Cluster_Label'
)['GPA'].mean().sort_values()

cluster_names = {
    cluster_means.index[0]: "Trung bình / Yếu",
    cluster_means.index[1]: "Khá",
    cluster_means.index[2]: "Giỏi"
}

features_filtered['Ten_Nhom_Cum'] = (
    features_filtered['Cluster_Label']
    .map(cluster_names)
)

print("\nThống kê số lượng:")

print(
    features_filtered['Ten_Nhom_Cum']
    .value_counts()
)

# ==============================================================================
# 10. BIỂU ĐỒ PHÂN CỤM
# ==============================================================================

np.random.seed(42)

plt.figure(figsize=(11,4))

colors_map = {
    "Giỏi": '#2E7D32',
    "Khá": '#FBC02D',
    "Trung bình / Yếu": '#C62828'
}

for group_name in [
    "Giỏi",
    "Khá",
    "Trung bình / Yếu"
]:

    sub_df = features_filtered[
        features_filtered['Ten_Nhom_Cum']
        == group_name
    ]

    y_jitter = np.random.normal(
        0,
        0.05,
        size=len(sub_df)
    )

    plt.scatter(
        sub_df['GPA'],
        y_jitter,
        c=colors_map[group_name],
        label=group_name,
        s=100,
        edgecolors='black',
        alpha=0.85
    )

plt.title(
    'Biểu đồ phân cụm sinh viên',
    fontsize=13,
    fontweight='bold'
)

plt.xlabel('Điểm GPA')

plt.xlim(1.0, 4.0)

plt.gca().get_yaxis().set_visible(False)

plt.grid(axis='x', linestyle=':')

plt.legend()

plt.tight_layout()

plt.savefig(
    'bieu_do_phan_cum.png',
    dpi=300
)

plt.show()

plt.close()

# ==============================================================================
# 11. BIỂU ĐỒ TRÒN
# ==============================================================================

labels = [
    "Giỏi",
    "Khá",
    "Trung bình / Yếu"
]

counts = [
    features_filtered[
        features_filtered['Ten_Nhom_Cum'] == l
    ].shape[0]
    for l in labels
]

plt.figure(figsize=(8,7))

plt.pie(
    counts,
    labels=labels,
    autopct='%1.1f%%',
    startangle=140
)

plt.title(
    'Tỷ lệ các cụm học lực'
)

plt.axis('equal')

plt.tight_layout()

plt.savefig(
    'bieu_do_tron.png',
    dpi=300
)

plt.show()

plt.close()

# ==============================================================================
# 12. XUẤT FILE CSV
# ==============================================================================

features_filtered.to_csv(
    "ket_qua_phan_cum.csv",
    index=False,
    encoding='utf-8-sig'
)

print("\nĐã xuất file CSV")

# ==============================================================================
# 13. SẮP XẾP DỮ LIỆU
# ==============================================================================

priority_map = {
    "Giỏi": 0,
    "Khá": 1,
    "Trung bình / Yếu": 2
}

features_filtered['Priority'] = (
    features_filtered['Ten_Nhom_Cum']
    .map(priority_map)
)

features_filtered = features_filtered.sort_values(
    by=['Priority', 'GPA'],
    ascending=[True, False]
).reset_index(drop=True)

# ==============================================================================
# 14. XUẤT FILE EXCEL ĐẸP
# ==============================================================================

wb = openpyxl.Workbook()

ws = wb.active

ws.title = "Phan Cum GPA"

# =====================
# STYLE
# =====================

HEADER_FILL = PatternFill(
    start_color="1F497D",
    end_color="1F497D",
    fill_type="solid"
)

FONT_HEADER = Font(
    bold=True,
    color="FFFFFF"
)

CENTER = Alignment(
    horizontal="center",
    vertical="center"
)

thin = Side(
    border_style="thin",
    color="CCCCCC"
)

border = Border(
    left=thin,
    right=thin,
    top=thin,
    bottom=thin
)

# =====================
# TITLE
# =====================

ws.merge_cells("A1:E1")

ws["A1"] = "DANH SACH PHAN CUM SINH VIEN"

ws["A1"].font = Font(
    bold=True,
    size=14
)

ws["A1"].alignment = CENTER

# =====================
# HEADER
# =====================

headers = [
    "STT",
    "MSSV",
    "Ten Sinh Vien",
    "GPA",
    "Nhom"
]

for col, header in enumerate(headers, 1):

    cell = ws.cell(
        row=3,
        column=col,
        value=header
    )

    cell.fill = HEADER_FILL

    cell.font = FONT_HEADER

    cell.alignment = CENTER

    cell.border = border

# =====================
# DATA
# =====================

row_num = 4

for idx, row in features_filtered.iterrows():

    data = [
        idx + 1,
        row['MSSV'],
        row['Ten_SV'],
        round(row['GPA'], 2),
        row['Ten_Nhom_Cum']
    ]

    for col, value in enumerate(data, 1):

        cell = ws.cell(
            row=row_num,
            column=col,
            value=value
        )

        cell.border = border

        if col != 3:
            cell.alignment = CENTER

    row_num += 1

# =====================
# WIDTH
# =====================

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 30
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 25

# =====================
# SAVE
# =====================

output_excel = "Ket_Qua_Phan_Cum.xlsx"

wb.save(output_excel)

print("\nĐã xuất file Excel")

print("\nHOÀN THÀNH!")